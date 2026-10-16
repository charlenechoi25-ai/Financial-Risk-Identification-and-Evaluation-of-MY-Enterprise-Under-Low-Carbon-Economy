"""Curate public evidence for chapter 3; retain conflicts and unknown cash receipts."""
from pathlib import Path
from decimal import Decimal as D
import csv
import hashlib
import json
import os
import re
import shutil

ROOT = Path(__file__).resolve().parents[1]
os.environ.setdefault('MPLCONFIGDIR', str(ROOT / '.cache/matplotlib'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill

OUT = ROOT / 'outputs/chapter3_support'
EXT = ROOT / 'data/raw/external'
DISC = EXT / 'discovery'
DATE = '2026-10-16'
SOURCES = {}
QUOTES = []
CHECKS = []


def read_csv(path):
    with (ROOT / path).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def norm(text):
    return re.sub(r'\s+', '', text)


def source(sid, stem, title, url, kind, published, folder=EXT, extension='html', doc_date=''):
    raw = folder / (stem + '.' + extension)
    txt = (ROOT / 'data/interim/pdf_text' if sid in ['A25','A24'] else folder) / (stem + '.txt')
    assert raw.exists() and txt.exists(), raw
    SOURCES[sid] = dict(source_id=sid, title=title, source_type=kind,
                        publication_date=published, document_date=doc_date,
                        access_date=DATE, url=url, local_path=str(raw.relative_to(ROOT)),
                        sha256=hashlib.sha256(raw.read_bytes()).hexdigest(),
                        text_sha256=hashlib.sha256(txt.read_bytes()).hexdigest(),
                        text_path=str(txt.relative_to(ROOT)))


def quote(qid, sid, text, meaning, limit, section='', page=None):
    pages = (ROOT / SOURCES[sid]['text_path']).read_text(encoding='utf-8').split('\f')
    matches = [i + 1 for i, p in enumerate(pages) if norm(text) in norm(p)]
    if page is not None:
        assert page in matches, (qid, page, text)
        matches = [page]
    assert matches, (qid, sid, text)
    is_pdf = SOURCES[sid]['local_path'].endswith('.pdf')
    QUOTES.append(dict(evidence_id=qid, source_id=sid,
                       locator=('PDF物理页' + ','.join(map(str, matches))) if is_pdf else section,
                       quote=text, interpretation=meaning, limit=limit,
                       verification='原文逐字核对（仅忽略空白）'))


def check(name, residual, expected='0', note=''):
    residual = D(str(residual))
    target = D(expected)
    assert abs(residual - target) <= D('.02'), (name, residual, target)
    CHECKS.append(dict(check=name, actual=str(residual), expected=str(target), passed=True, note=note))


def register_sources():
    source('P01','carbon_action_2030','2030年前碳达峰行动方案','https://www.gov.cn/zhengce/content/2021-10/26/content_5644984.htm','政府政策原文','2021-10-26',doc_date='2021-10-24')
    source('P02','renewable_pricing_136','新能源上网电价市场化改革136号文','https://www.ndrc.gov.cn/xxgk/zcfb/tz/202502/t20250209_1396066.html','政府政策原文','2025-02-09',doc_date='2025-01-27')
    source('P03','wind_price_882','关于完善风电上网电价政策的通知882号文','https://www.ndrc.gov.cn/xxgk/zcfb/tz/201905/t20190524_962453.html','政府政策原文','2019-05-24',doc_date='2019-05-21')
    source('P04','subsidy_4','关于促进非水可再生能源发电健康发展的若干意见','https://www.gov.cn/zhengce/zhengceku/2020-02/03/content_5474144.htm','政府政策原文','2020-02-03',doc_date='2020-01-20')
    source('P05','parity_833','关于2021年新能源上网电价政策有关事项的通知','https://www.ndrc.gov.cn/xxgk/zcfb/tz/202106/t20210611_1283088.html','政府政策原文','2021-06-11',doc_date='2021-06-07')
    urls = {
        2019: ('2020-02-28','https://www.nea.gov.cn/2020-02/28/c_138827910.htm','wind2019'),
        2020: ('2021-01-30','https://www.nea.gov.cn/2021-01/30/c_139708580.htm','wind2020'),
        2021: ('2022-01-28','https://www.nea.gov.cn/2022-01/28/c_1310445390.htm','wind2021'),
        2022: ('2023-02-13','https://www.nea.gov.cn/2023-02/13/c_1310697149.htm','wind2022'),
        2023: ('2024-01-26','https://www.nea.gov.cn/2024-01/26/c_1310762246.htm','power2023'),
        2024: ('2025-01-27','https://www.nea.gov.cn/20250221/e10f363cabe3458aaf78ba4558970054/c.html','wind2024'),
        2025: ('2026-02-12','https://www.nea.gov.cn/20260212/742b8c6a078347b0b39de676c05c5d58/c.html','wind2025'),
    }
    for year, (pub, url, stem) in urls.items():
        source('I'+str(year),stem,f'{year}年风电官方运行/电力统计资料',url,'国家能源局统计原始发布',pub)
    source('J24','wind_industry_2024','中国风电产业持续迭代升级报道','https://www.nea.gov.cn/2024-10/18/c_1310786778.htm','能源局转载央视网报道（非官方统计原始发布）','2024-10-18')
    source('J25','wind_industry_2025','风电行业从规模扩张转向价值提升','https://www.nea.gov.cn/20251107/3f55745bfb544028bdf8bad79615882d/c.html','能源局转载科技日报（含受访企业观点）','2025-11-07')
    docs = [
        ('TK','transfer_20240622','AN202406211636768053','开鲁明阳股权转让公告2024-061','2024-06-22'),
        ('TN','transfer_20240626','AN202406251636874015','奈曼旗明阳股权转让公告2024-063','2024-06-26'),
        ('TP','transfer_progress_20241001','AN202409301640102972','股权转让进展公告2024-087','2024-10-01'),
        ('TE','transfer_explanation_20241010','AN202410091640229634','股权转让进展说明2024-089','2024-10-10'),
        ('TT','taonan2025','AN202503261647797407','洮南百强股权转让公告2025-015','2025-03-27'),
        ('R24','company_reply_2024','AN202407111637805784','公司2023年年报监管工作函回复2024-067','2024-07-12'),
        ('RA24','regulator_reply_2024','AN202407111637805782','致同专字（2024）第110A015118号回复','2024-07-12'),
        ('R23','regulator_reply_2023','AN202307121592217284','公司2022年年报监管工作函回复2023-043','2023-07-13'),
        ('H26','interim_2026','AN202608281828646703','2026年半年度报告（未经审计，仅期后观察）','2026-08-29'),
        ('C25','credit_2025','AN202504251662324475','2025年度拟申请授信额度公告2025-026','2025-04-26'),
    ]
    for sid, stem, code, title, pub in docs:
        kind='会计师回复文件（含公司说明）；东方财富镜像' if sid=='RA24' else '发行人披露文件；东方财富镜像（不是交易所托管URL）'
        source(sid,stem,title,f'https://pdf.dfcfw.com/pdf/H2_{code}_1.pdf',kind,pub,DISC,'pdf')
    for sid, stem, year in [('A25','20260428',2025),('A24','20250425',2024)]:
        source(sid,stem,f'{year}年年度报告（用户上传版本）','用户上传文件，参照来源路径和SHA256','用户提供年报原件','',ROOT/'data/raw/uploads','pdf')
        # Uploaded text lives separately; keep the PDF itself unmodified.


def collect_quotes():
    quote('E01','P03','2018年底之前核准的陆上风电项目，2020年底前仍未完成并网的，国家不再补贴','区分核准批次和并网窗口','不将陆上项目规则套用于全部海上项目','一（四）')
    quote('E02','P03','2019年1月1日至2020年底前核准的陆上风电项目，2021年底前仍未完成并网的，国家不再补贴','另一核准批次的并网窗口','不能简化为全部陆上风电2020年统一退补','一（四）')
    quote('E03','P04','新增海上风电和光热项目不再纳入中央财政补贴范围','中央财政新增项目范围变化','不代表地方支持政策一并取消','一（一）')
    quote('E04','P04','并于2021年12月31日前全部机组完成并网的存量海上风力发电和太阳能光热发电项目','存量海上项目有核准及全部并网等条件','不能只凭单个日期认定某项目补贴资格','一（一）')
    quote('E05','P05','2021年起，对新备案集中式光伏电站、工商业分布式光伏项目和新核准陆上风电项目','平价政策明确适用对象','不覆盖全部历史存量项目','一')
    quote('E06','P05','本通知自2021年8月1日起执行。','明确执行日','成文日、网页发布日期与执行日分别列示','末段')
    quote('E07','P01','全面推进风电、太阳能发电大规模开发和高质量发展','清洁能源产业需求背景','国家方向不等于公司实际订单或收益','三（一）2')
    quote('E08','P02','市场交易均价低于或高于机制电价的部分，由电网企业按规定开展差价结算','市场交易与缓释机制并存','仅适用于纳入机制的电量，不能推断收益必然下降','三（四）')
    quote('E09','P02','2025年6月1日起投产的新能源增量项目','按投产日期区分规则','不是按公司股权处置日划分，不能追溯解释2024年调价','三（五）')
    quote('E10','P02','各地要在2025年底前出台并实施具体方案','地方实施有差异','本轮未把全国文件当作所有项目结算参数','末段')
    quote('E11','J24','从陆上大功率到深耕深远海','装备技术迭代背景','转载新闻是行业背景，不等于公司技术失败或统计原始发布','正文')
    quote('E12','J25','风电行业从规模扩张转向价值提升','竞争从规模向价值演进的报道线索','含采访及企业观点，不引用其全保底退出等概括替代136号文条款','标题')
    quote('E13','TE','基于消纳、电价等市场及行业政策环境变化，公司通过价格调整推进两项目的出售','企业明确披露的环境与交易价格关联','属于公司解释，不是独立识别的因果效应，更不是2025年136号文的影响',page=1)
    quote('E14','TP','81,623.86万元和48,678.94万元','两项目最新公告转让价格','不能将原始价款仍作为最新价格',page=1)
    quote('E15','TP','第三期、第四期、第五期股权转让款不再执行','原约定远期三期被取消','其他支付金额与新总价的公开资料算术衔接存在差额，勿强行配平',page=2)
    quote('E16','TE','其中2.2亿元为《委托贷款借款合同》剩余未偿还本金','项目借款与股权价款分开','不是总转让价款的一部分，未核实实际到账',page=2)
    quote('E17','TE','目前公司对两个项目的担保已解除','处置同时降低担保敞口的缓释证据','不能把旧担保继续列作期末未解除事项',page=2)
    quote('E18','TT','在消缺保证期届满前，未有影响项目运营事项并完成消缺事项后支付5%','尾款附有具体经营条件','条件支付不自动等于逾期',page=5)
    quote('E19','TT','该笔费用在第三笔股权对价转让款中扣减','146万元利息负担影响尾款现金安排','按比例推算额不等于最终实际收款',page=6)
    quote('E20','TT','洮南百强应付北京洁源11,100万元','借款/往来回收安排与股权款分别观察','不直接与2023年工程应收或合同资产连成同一科目余额',page=6)
    quote('E21','TT','工商变更登记完成之日起20个工作日内','该1.11亿元有单独支付时点安排','工商变更日和付款凭证未取得，不能自行计算已逾期',page=6)
    quote('E22','R24','升压站消防水池、消防水泵等设施需完善','洮南百强合同资产有具体消缺事项','不能将个案原因推广至2025年全部合同资产')
    quote('E23','R24','截至2024年6月30日已转至应收账款','结算条件满足后的科目转化','转应收不是现金收回')
    quote('E24','R24','叶片震动监测系统接入尚未达到预期状态','另一EPC项目的结算条件','当时预计完成时间不是已完成事实')
    quote('E25','R24','尚未达到合同结算的时间节点，待质保期结束后可结转至应收账款','依安明阳质保金形成合同资产的个案','不能据此认定全部合同资产都是质保金')
    quote('E26','R24','一般需要1-3年','当时披露的预验收周期说明','是该回复中的业务说明，不是所有2025年合同的刚性期限')
    quote('E27','R24','预付款、投料款、到货款、预验收款和质保金','制造回款节点与收入确认有差异','不把收入确认、取得收款权和实际到账混为一谈')
    quote('E28','R24','短期的运营资金不用于长期投资','公司解释的资金分池管理','属于管理层说明，不等于对2025年所有借款用途独立验证')
    quote('E29','C25','最终授信额度将以公司及子公司与金融机构签订的合同为准','拟申请授信不等于已获批或未使用额度','不以拟申请1370.01亿元充当可动用备用现金',page=4)
    quote('E30','H26','本半年度报告未经审计。','期后材料审计状态','只作2025年末风险信号跟踪，不混入2020—2025年度评价样本',page=2)
    quote('E31','R23','截至2023年5月31日已收到期后回款18.24亿元','2022年末应收的期后实际回款披露','期间为2023年1—5月，不能当作2025年的回款表现',page=17)
    quote('E32','R24','2,426.20','2023公司平均千瓦收入，与2022年3,121.46对照','公司指标，不是全行业招标均价')
    quote('E33','A25','截至本报告披露日，公司尚未就上述项目投入资金','苏格兰基地仍为规划而非已支付投资','142.10亿元折算计划额不能加入2025年实际资本支出',page=32)
    quote('E34','A25','产品质量保证','预计负债及审计估计的重要性','不能仅凭余额或关键审计事项判断质量事故或计提不足',page=193)
    quote('E35','A25','本报告期末已签订合同','剩余履约义务的披露','不是全部市场订单，更不是已收现金',page=200)
    quote('E36','A25','2006年6月2日成立','企业成立日期','以年报基本情况为准，不将股改或上市日当成立日',page=115)
    quote('E37','A25','2019年1月23日挂牌上市交易','A股上市日期','不与GDR发行混同',page=115)
    quote('E38','TT','公司此前为洮南百强提供的连带责任担保已解除','项目担保缓释证据','本公告同时披露本次交易尚需债权人同意，保持时点',page=1)
    quote('E39','TT','已建成投运','洮南百强为成熟运营项目','不写作仍未建成的失败项目',page=3)


def industry():
    # National installed/grid-connected capacity; never mix with lifted/tender capacity.
    vals = [(2019,'2.1亿千瓦','210','全国风电累计装机2.1亿千瓦','2574'),
            (2020,'2.81亿千瓦','281','全国风电累计装机2.81亿千瓦','7167'),
            (2021,'3.28亿千瓦','328','全国风电累计装机3.28亿千瓦','4757'),
            (2022,'3.65亿千瓦','365','风电3.65亿千瓦','3763'),
            (2023,'44134万千瓦','441.34','风电万千瓦4413420.7',''),
            (2024,'5.21亿千瓦','521','全国风电累计并网容量达到5.21亿千瓦','7982'),
            (2025,'6.4亿千瓦','640','全国风电累计并网容量达到6.4亿千瓦','12000')]
    rows=[]
    for y, original, gw, literal, added in vals:
        sid='I'+str(y)
        quote('I'+str(y),sid,literal,'全国风电累计装机/并网规模','保持各年度原始精度；不是吊装或招标容量','风电段/统计表')
        if added:
            # 2025 is disclosed as 1.2 亿千瓦, not a precise 12000.00 万千瓦 observation.
            token='1.2亿千瓦' if y==2025 else added+'万千瓦'
            assert token in norm((ROOT/SOURCES[sid]['text_path']).read_text())
        rows.append(dict(year=y,capacity_original=original,capacity_GW=gw,
                         new_capacity_10k_kW=added,source_id=sid,
                         definition='全国风电装机/并网统计；非吊装、非招标',
                         precision_note='2023取统计表精确值；其余按正文原有舍入值，不据此重算精确同比',
                         missing_note='本轮未取得2023新增并网原始统计，不用存量差额或媒体75GW补齐' if not added else ('2025原披露1.2亿千瓦为舍入值' if y==2025 else '')))
    return rows


def projects():
    revisions=[]
    for name, old, new in [('开鲁县明阳智慧能源有限公司','959902700','816238600'),('奈曼旗明阳智慧能源有限公司','557511000','486789400')]:
        old,new=D(old),D(new)
        revisions.append(dict(project=name,original_price_cny=str(old),revised_price_cny=str(new),
                              reduction_cny=str(old-new),reduction_pct=str((old-new)/old*100),
                              source_ids='TK/TN;TP;TE',cause='公司公告归因于消纳、电价等市场及行业政策环境变化',
                              remaining_85pct_original=str(old*D('.85')),
                              unreconciled_difference_cny=str(new-old*D('.85')),
                              status='公告新总价已核验；旧比例与保留金额未逐项桥接',actual_cash_received_cny=None))
    check('two_project_price_reduction',sum(D(r['reduction_cny']) for r in revisions),'214385700')
    payments=[]
    for sid, project in [('TK',revisions[0]['project']),('TN',revisions[1]['project'])]:
        for n,pct,condition in [(1,51,'董监高变更、章证照移交、资产交割'),(2,24,'过渡期审计、股权工商变更'),(3,5,'原约定满两年'),(4,5,'原约定满四年'),(5,5,'原约定满六年'),(6,10,'完成消缺')]:
            payments.append(dict(project=project,stage=n,original_percentage=pct,condition=condition,latest_status='2024补充协议取消' if n in [3,4,5] else '公告称金额、条件不变；精确金额桥接未完成',nominal_original_price_calculation_cny=None,actual_cash_received_cny=None,source_ids=sid+';TP'))
    taonan=D('96830100')
    for n,pct,condition in [(1,51,'集团备案、相关协议、资料移交及债权人确认'),(2,44,'国资产权登记、风机权利义务转移、质保保函及工商变更'),(3,5,'消缺保证期内无影响运营事项并完成消缺')]:
        payments.append(dict(project='洮南百强新能源有限公司',stage=n,original_percentage=pct,condition=condition,latest_status='2025-015合同安排；第3期另扣146万元利息',nominal_original_price_calculation_cny=str(taonan*D(pct)/100),actual_cash_received_cny=None,source_ids='TT'))
    check('taonan_three_nominal_instalments',sum(D(r['nominal_original_price_calculation_cny']) for r in payments if r['source_ids']=='TT')-taonan)
    tracking=[]
    for company, old, new, kind in [('中广核风电有限公司','207890000','168490000','股权转让款'),('锡林郭勒盟恒阳新能源有限公司','126379088.25','126379088.25','往来款'),('湖南国核新能源有限公司','67261756.12','67261756.12','股权转让款')]:
        assert format(D(old),',.2f') in (ROOT/SOURCES['A25']['text_path']).read_text().split('\f')[156]
        assert format(D(new),',.2f') in (ROOT/SOURCES['H26']['text_path']).read_text().split('\f')[103]
        tracking.append(dict(counterparty=company,nature=kind,balance_2025_12_31_cny=old,balance_2026_06_30_cny=new,net_decrease_cny=str(D(old)-D(new)),actual_receipt_by_counterparty_cny=None,source_ids='A25 PDF157;H26 PDF104',status='2026半年报未经审计，期后辅助观察',limit='余额差额不是已核验收现；中广核对手一致但不能唯一分配到开鲁/奈曼旗具体项目'))
    check('cgn_balance_decrease',D(tracking[0]['net_decrease_cny']),'39400000')
    check('subsequent_aggregate_cash_amount_equal_to_cgn_change',D('39400000')-D(tracking[0]['net_decrease_cny']),note='仅金额吻合；不标记为对手或项目现金匹配成功')
    return revisions,payments,tracking


def financial_indicators():
    rows=read_csv('data/processed/financial_statement_versions.csv')
    index={(r['report_year'],r['version'],r['metric']):r for r in rows}
    result=[]
    for year in range(2019,2026):
        used=[]
        def value(metric,version='current'):
            r=index[(str(year),version,metric)]
            assert r['value']!='', (year,metric,version)
            used.append(r)
            return D(r['value'])
        a,l,ca,cl=value('total_assets'),value('total_liabilities'),value('current_assets'),value('current_liabilities')
        rev,profit,cfo,cost=value('revenue'),value('net_profit'),value('cfo'),value('cost')
        ar,ar0=value('receivables'),value('receivables','comparative')
        inv,inv0=value('inventory'),value('inventory','comparative')
        ratios=[('资产负债率',l/a*100,'%', '负债合计/资产总计×100'),('流动比率',ca/cl,'倍','流动资产/流动负债'),('销售净利率',profit/rev*100,'%','合并净利润/营业收入×100'),('经营现金流收入比',cfo/rev*100,'%','经营净现金流/营业收入×100'),('应收账款周转天数',((ar+ar0)/2)/rev*365,'天','同份年报期初期末应收账款净额均值/营业收入×365'),('存货周转天数',((inv+inv0)/2)/cost*365,'天','同份年报期初期末存货净额均值/营业成本×365')]
        for name,val,unit,formula in ratios:
            result.append(dict(year=year,indicator=name,value=str(val),unit=unit,formula=formula,
                               source_file=used[0]['source_file'],pdf_pages=','.join(sorted({r['pdf_page'] for r in used},key=int)),
                               basis='每年本报告期披露；周转均值取同份年报比较数，跨年未完全重述统一',
                               limit='描述性识别，不含评分、阈值或行业风险等级；周转天数不是合同逾期天数'))
        check(f'{year}_asset_identity',a-l-value('total_equity'))
    return result


def supplemental_observations():
    records=[
        ('A25',24,'2025','风机及相关配件销售收入','31,010,004,661.15','元','分产品主营业务'),
        ('A25',24,'2025','电站产品销售收入','4,506,462,857.05','元','分产品主营业务'),
        ('A25',24,'2025','风电场发电收入','1,353,690,343.84','元','分产品主营业务'),
        ('A25',24,'2025','风机及相关配件毛利率','6.38','%','同比提高6.84个百分点'),
        ('A25',24,'2025','电站产品销售毛利率','28.99','%','同比降低14.18个百分点'),
        ('A25',24,'2025','风电场发电毛利率','45.99','%','同比降低11.15个百分点'),
        ('A25',200,'2025末','剩余履约义务对应收入','11,921,755,195.08','元','不是全部订单或已收现金'),
        ('A25',190,'2025末','一年内到期长期借款','1,163,762,581.72','元','区别含预计负债的全部一年内到期非流动负债'),
        ('A25',190,'2025末','一年内到期预计负债','407,405,979.04','元','不得全当有息借款'),
        ('A25',202,'2025','财务费用','515,530,966.30','元','不能仅由短借增长解释'),
        ('A25',202,'2024比较','财务费用','342,440,213.42','元','同表比较'),
        ('A25',202,'2025','利息支出','485,190,370.76','元','另有资本化、利息收入、汇兑等项目'),
        ('A25',202,'2024比较','利息支出','528,233,235.58','元','利息支出下降与财务费用上升并存'),
        ('A25',193,'2025末','产品质量保证预计负债','1,731,062,057.99','元','余额不是本期实际质保现金支出'),
        ('A25',193,'2024末比较','产品质量保证预计负债','1,759,258,369.32','元','余额下降不能自动等同质保支出减少'),
        ('A25',214,'2025末','8-10MW项目预计完成日','2026/7/31','日期','披露时为预期；不是已完成日'),
        ('A25',214,'2025末','16-18MW项目预计完成日','2026/9/30','日期','披露时为预期；不是已完成日'),
        ('H26',159,'2026半年报期后','8-10MW项目预计完成日','2027/5/31','日期','期后预期调整；未披露原因或实际失败'),
        ('H26',159,'2026半年报期后','16-18MW项目预计完成日','2026/12/31','日期','期后预期调整；不能倒填2025已知信息'),
        ('H26',154,'2026上半年期后','以前期间处置收到的现金','39,400,000.00','元','汇总收现事实，逐项目/逐对手归属未核实'),
        ('H26',104,'2026-06-30期后','中广核风电股权转让款余额','168,490,000.00','元','余额净减少不自动等于该对手实际收现'),
        ('H26',104,'2026-06-30期后','恒阳往来款余额','126,379,088.25','元','余额未变不证明期间无新增或无回款'),
        ('H26',104,'2026-06-30期后','湖南国核股权转让款余额','67,261,756.12','元','无合同到期安排，不直接判逾期'),
        ('C25',1,'2025计划','拟申请授信总额','13,700,100','万元','1370.01亿元；不是获批、已使用或未使用授信'),
        ('TT',6,'2025公告安排','洮南百强应付北京洁源','11,100','万元','与股权价款分开；约定付款日不等于实际付款'),
        ('TT',6,'2025公告安排','北京洁源承担利息','146','万元','在第三期股权价款中扣减'),
    ]
    result=[]
    for sid,page,period,label,val,unit,limit in records:
        text=(ROOT/SOURCES[sid]['text_path']).read_text().split('\f')[page-1]
        assert norm(val) in norm(text),(sid,page,label,val)
        result.append(dict(source_id=sid,pdf_page=page,period=period,item=label,original_value=val,unit=unit,limit=limit))
    return result


def finance_bridge():
    items=[('利息支出','485190370.76','528233235.58',1),('未确认融资费用转回','117251120.04','122235983.36',1),('减：利息资本化','69570656.84','144207507.67',-1),('减：利息收入','137080096.04','203650412.13',-1),('汇兑损益','45868624.07','-24488838.69',1),('手续费及其他','73871604.31','64317752.97',1)]
    text=(ROOT/SOURCES['A25']['text_path']).read_text().split('\f')[201]
    result=[]
    for label,current,prior,sign in items:
        for v in [current,prior]:assert format(D(v),',.2f') in text,(label,v)
        result.append(dict(item=label,amount_2025_cny=current,amount_2024_cny=prior,sign=sign,expense_change_cny=str((D(current)-D(prior))*sign),source='A25 PDF202',limit='财务费用的会计项目调节，不等同全部借款现金利息'))
    check('finance_expense_2025_components',sum(D(r['amount_2025_cny'])*r['sign'] for r in result),'515530966.30')
    check('finance_expense_2024_components',sum(D(r['amount_2024_cny'])*r['sign'] for r in result),'342440213.42')
    return result


def write_outputs(sheets):
    wb=Workbook();wb.remove(wb.active)
    for name,rows in sheets.items():
        assert rows,name
        fields=list(dict.fromkeys(k for r in rows for k in r))
        with (OUT/(name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
        ws=wb.create_sheet(name[:31]);ws.append(fields)
        for r in rows: ws.append([r.get(k) for k in fields])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='245269')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=32
            for c in col:c.alignment=Alignment(vertical='top',wrap_text=True)
        for i in range(2,ws.max_row+1):ws.row_dimensions[i].height=72
    path=OUT/'第三章完整写作支撑证据.xlsx';wb.save(path)
    with_check=load_workbook(path,read_only=True)
    assert len(with_check.sheetnames)==len(sheets)
    for ws in with_check:
        for row in ws.iter_rows(values_only=True):pass
    with_check.close()


def plots(industry_rows,revisions):
    font=FontProperties(fname='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    plt.rcParams['font.family']=font.get_name();plt.rcParams['axes.unicode_minus']=False
    fig,ax=plt.subplots(figsize=(9,4.8));ys=[float(r['capacity_GW']) for r in industry_rows]
    ax.plot([r['year'] for r in industry_rows],ys,marker='o',color='#237a80')
    for r,y in zip(industry_rows,ys):ax.annotate(r['capacity_GW'],(r['year'],y),xytext=(0,8),textcoords='offset points',ha='center')
    ax.set(title='全国风电装机规模：2019—2025年',ylabel='GW（吉瓦）',xlabel='统计年度')
    ax.set_ylim(min(ys)*.9,max(ys)*1.12)
    ax.grid(alpha=.2);fig.text(.03,.02,'来源：国家能源局原始发布；各年原披露精度不同。非吊装/招标容量；不由存量差额推算新增。',fontsize=9)
    fig.tight_layout(rect=(0,.06,1,1));fig.savefig(OUT/'全国风电装机规模.png',dpi=180);plt.close(fig)
    fig,ax=plt.subplots(figsize=(8,4.8));x=[0,1]
    ax.bar([i-.18 for i in x],[float(D(r['original_price_cny'])/D('1e8')) for r in revisions],width=.36,label='2024年6月原公告价款')
    ax.bar([i+.18 for i in x],[float(D(r['revised_price_cny'])/D('1e8')) for r in revisions],width=.36,label='2024年补充协议公告价款')
    ax.set_xticks(x,['开鲁明阳','奈曼旗明阳']);ax.set(ylabel='股权转让价款（亿元）',title='两电站项目转让价款调整');ax.legend();ax.grid(axis='y',alpha=.2)
    fig.text(.04,.02,'来源：2024-061、063、087、089公告；调整差额不等于当期投资损失或现金损失。',fontsize=9)
    fig.tight_layout(rect=(0,.06,1,1));fig.savefig(OUT/'电站项目转让价款调整.png',dpi=180);plt.close(fig)


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    SOURCES.clear();QUOTES.clear();CHECKS.clear()
    register_sources();collect_quotes();ir=industry();revisions,payments,tracking=projects();indicators=financial_indicators();observations=supplemental_observations();finance=finance_bridge()
    conflicts=[
        dict(id='D01',topic='两项目调价与原分期比例',source='TK/TN;TP',issue='按旧价51%+24%+10%测算保留价款，与最新总价分别差321305元和12905050元',treatment='采用最新公告总价；不将原比例乘新总价，也不强制配平各期；不估计实际回款',blocking='不影响调价事实；阻止精确付款计划重建'),
        dict(id='D02',topic='洮南百强2023年末工程应收是否在信用期',source='R24 PDF15与27；RA24 PDF17',issue='工程汇总表将含应收及合同资产的往来列为信用期内；公司后续应收细表列5589.81万元为信用期外',treatment='两处表格原样并列，不将整笔合同资产认作逾期；细表标示的信用期外信号单独披露并注明口径分歧',blocking='阻止对该项目全余额作统一逾期判断'),
        dict(id='D03',topic='2026期后中广核余额变化与现金',source='A25 PDF157;H26 PDF104/154',issue='对手余额净减少3940万元，与汇总以前期间处置收现金额3940万元相等',treatment='仅列金额吻合待匹配；未取得逐项目收款归属，不宣称项目现金匹配完成',blocking='阻止按项目分摊实际收款'),
        dict(id='D04',topic='风电利用率跨年基数',source='I2020;I2021',issue='2020原披露97%，2021原披露96.9%但称同比提高0.4个百分点，存在前期比较基数差异',treatment='不拼接成统一口径连续利用率曲线；仅就2024的95.9%与2025的94%作原披露比较，并提示方法口径待核实',blocking='不影响装机规模序列；阻止强行解释全期利用率同比'),
        dict(id='D05',topic='未折现到期负债与借款利率敏感性',source='A25 PDF233',issue='2024与2025到期分类列示不同；2025利率敏感性方向与常见直觉及2024列示不一致',treatment='不将该表直接拼成统一借款本金序列，不据此测算利率冲击；保留待核验',blocking='不影响BS借款结构与资金需求描述'),
    ]
    readiness=[]
    for section,status,evidence,limit in [
        ('3.1.1 低碳经济发展与风电产业政策','可写','P01—P05；E01—E10','按历史阶段和项目适用范围表述，不将2025政策追溯解释2024调价'),
        ('3.1.2 市场竞争与技术变化','可写（限定证据范围）','I2019—I2025；J24/J25；R24；A25','公司售价不是行业招标价，媒体观点不是官方统计；不声称建立完整招标价格序列'),
        ('3.1.3 行业条件与财务传导','可写','TE实际调价案例；TT付款条件；R24结算事项','案例关联和公司归因不等于统计因果检验'),
        ('3.2.1 企业基本情况','可写','A25 PDF6—7及115公司基本情况','沿用论文MY匿名；不混入2026控制权变化作为2025事实'),
        ('3.2.2 主营业务与经营特征','可写','A25 PDF12—13/24/124/131/200；R24','区分制造、电站产品、自持发电及不同资产会计归属'),
        ('3.2.3 财务状况分析','可写（披露版本说明必留）','2019—2025原披露值、42个描述性指标、版本核对表','不声称跨年已完全重述统一；周转天数不是逾期'),
        ('3.3.1 筹资风险信号','可写（期限匹配问题）','BS借款、A25 PDF190/202/232；R24 PDF52—53；C25','拟申请授信不等于可用授信，四年度现金差额不用于覆盖率结论'),
        ('3.3.2 投资风险信号','可写（项目估值与回收）','TK/TN/TP/TE/TT；A25 PDF220；H26期后独列','不把减价等同损失，不编造逐项目到账，投资计划不当作已投入'),
        ('3.3.3 营运风险信号','可写（证据最充分）','既有05/06；R24结算条件、信用期外细表；期后跟踪','保留信用期披露分歧；不能把全部应收或合同资产判作逾期'),
        ('3.3.4 发展风险信号','可写（投入与转化不确定性）','研发投入、A25 PDF214、H26 PDF159，毛利及质保反证','预计完成日调整为期后观察，不倒填2025已知信息；不认定研发失败'),
        ('3.4 本章小结','可写','四维结论—证据—反证矩阵','仅识别风险信号，第四章权重/评分/等级仍未开展'),
    ]:readiness.append(dict(section=section,status=status,evidence=evidence,limit=limit))
    policies=[]
    for sid,effective,scope in [('P03','2019-07-01','陆上按核准批次和并网时间区分，海上另列'),('P04','未把成文日自动当生效日','新增海上中央补贴范围变化；符合条件的存量按并网窗口衔接'),('P05','2021-08-01','2021起新核准陆上平价；海上省级定价等规定'),('P01','2025/2030目标期；非统一生效日期','清洁能源规模发展及高质量发展方向'),('P02','2025-06-01为投产分界，地方2025年底前实施','存量/增量分类；交易价格和机制电量差价结算并存')]:
        s=SOURCES[sid];policies.append(dict(source_id=sid,title=s['title'],document_date=s['document_date'],publication_date=s['publication_date'],execution_or_boundary=effective,scope=scope,url=s['url']))
    guide=[dict(item='使用顺序',note='先读章节写作支撑稿和准备度，再查证据ID、来源目录及原文定位；沿用05/06数据基础，07为上阶段政策快照。'),dict(item='年份',note='主体为2020—2025，2019作基期；2026半年报仅作未经审计的期后跟踪。'),dict(item='数字',note='原始财务数值保留元；正文表用亿元、百分比；空白实际收款不是零。'),dict(item='结论',note='可写风险信号，不代表已完成第四章模型或第五章系统文本识别。'),dict(item='来源',note='政策与统计原文独立核验；发行人/会计师文件经公开镜像获取；媒体转载单独标记。')]
    sheets={'阅读指南':guide,'来源目录':list(SOURCES.values()),'原文证据索引':QUOTES,'政策时间线':policies,'行业装机统计':ir,'项目价款调整':revisions,'项目付款条件':payments,'期后回收观察':tracking,'补充观测与时点':observations,'财务费用变动桥接':finance,'描述性财务指标':indicators,'原披露财务序列':read_csv('outputs/chapter3/annual_financial_signals.csv'),'研发投入序列':read_csv('outputs/chapter3/rd_inputs.csv'),'争议与不可推断项':conflicts,'章节准备度':readiness,'计算核验':CHECKS}
    write_outputs(sheets);plots(ir,revisions)
    inputs={}
    for p in [ROOT/'data/processed/financial_statement_versions.csv',ROOT/'outputs/chapter3/annual_financial_signals.csv',ROOT/'outputs/chapter3/rd_inputs.csv']:
        inputs[str(p.relative_to(ROOT))]=hashlib.sha256(p.read_bytes()).hexdigest()
    manifest=dict(date=DATE,script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),source_count=len(SOURCES),quote_count=len(QUOTES),numeric_and_date_observations_checked=len(observations),quote_checks='exact normalized text; whitespace only ignored',indicator_count=len(indicators),checks=len(CHECKS),unresolved_conflicts=len(conflicts),project_level_actual_cash_matches=0,sources=list(SOURCES.values()),inputs=inputs)
    (OUT/'run_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:v for k,v in manifest.items() if k not in ['sources','inputs']},ensure_ascii=False))


if __name__=='__main__':main()
